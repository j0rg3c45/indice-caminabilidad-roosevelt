"""
Agrega la sección "Índice de Caminabilidad" al informe Word existente de Barrio Obrero,
SIN regenerar el documento (para conservar las ediciones del usuario, p. ej. la fuente Arial).

Respeta el formato del documento: fuente Arial y tamaño de cuerpo 11 pt, y usa los mismos
tamaños/colores de encabezado que el resto del informe.

Uso:
    uv run notebooks_py/agregar_indice_a_informe_docx.py
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

PROJECT_ROOT = Path(__file__).resolve().parent.parent
FIG_DIR = PROJECT_ROOT / "outputs" / "figures"
DOCX = PROJECT_ROOT / "outputs" / "informes" / "informe_caminabilidad_barrio_obrero.docx"

AZUL = RGBColor(0x1B, 0x4F, 0x8A)
VERDE = RGBColor(0x2E, 0x7D, 0x32)
GRIS = RGBColor(0x44, 0x44, 0x44)

FUENTE = "Arial"          # fuente del documento del usuario
SIZE_CUERPO = Pt(11)      # tamaño de cuerpo (docDefault = 22 half-points = 11 pt)
SIZE_H1 = Pt(14)          # Heading 1 del documento
SIZE_H2 = Pt(13)          # Heading 2 del documento

# Datos del índice (línea base)
INDICE = [
    ("Indicador", "ref_min", "ref_max", "Sentido", "Score BO", "Score C9"),
    ("Densidad de intersecciones (int/km²)", "100", "800", "directo", "73.9", "96.7"),
    ("Densidad de calle (km/km²)", "20", "90", "directo", "73.4", "83.4"),
    ("Longitud prom. de segmento (m)", "40", "150", "inverso", "98.6", "100.0"),
    ("Índice de Caminabilidad (0–100)", "", "", "", "78.7", "92.7"),
]


def _run(par, texto, size=SIZE_CUERPO, bold=False, color=None):
    r = par.add_run(texto)
    r.font.name = FUENTE
    r.font.size = size
    r.bold = bold
    if color is not None:
        r.font.color.rgb = color
    return r


def _parrafo(doc, texto, size=SIZE_CUERPO, bold=False, color=None, justify=True, space_after=8):
    p = doc.add_paragraph()
    _run(p, texto, size=size, bold=bold, color=color)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _titulo(doc, texto, size, color):
    p = doc.add_paragraph()
    _run(p, texto, size=size, bold=True, color=color)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def main():
    if not DOCX.exists():
        raise SystemExit(f"No existe el informe: {DOCX}. Genéralo primero.")
    doc = Document(DOCX)

    # Título de sección (nivel 1)
    _titulo(doc, "Índice de Caminabilidad (0–100)", SIZE_H1, AZUL)

    _parrafo(doc,
             "Para resumir la caminabilidad en una sola cifra, construimos un índice de 0 a "
             "100 que combina tres indicadores de la red peatonal. La normalización usa "
             "referencias fijas por indicador (no valores relativos entre zonas), de modo que "
             "el resultado sea estable y comparable a lo largo del tiempo. Cada indicador se "
             "lleva a una escala de 0 a 100 y luego se pondera; la longitud de segmento se "
             "trata como indicador inverso, ya que las cuadras más cortas favorecen la "
             "caminabilidad.")

    _parrafo(doc,
             "Los pesos por defecto son: densidad de intersecciones 0.45, densidad de calle "
             "0.35 y longitud de segmento 0.20. Tanto los umbrales como los pesos pueden "
             "ajustarse según el criterio del equipo. La siguiente tabla resume las "
             "referencias empleadas y los puntajes obtenidos por cada zona.")

    # Tabla del índice
    tabla = doc.add_table(rows=1, cols=6)
    tabla.style = "Light Grid Accent 1"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tabla.rows[0].cells
    for j, titulo in enumerate(INDICE[0]):
        hdr[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(hdr[j].paragraphs[0], titulo, size=Pt(9), bold=True)
    for fila in INDICE[1:]:
        row = tabla.add_row().cells
        es_total = fila[0].startswith("Índice")
        for j, val in enumerate(fila):
            row[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(row[j].paragraphs[0], val, size=Pt(9), bold=es_total)

    _parrafo(doc,
             "En la línea base, el Barrio Obrero obtiene un índice de 78.7 y la Comuna 9 de "
             "92.7. Ambos valores son altos y reflejan tramas urbanas caminables. La ventaja "
             "de la Comuna 9 proviene sobre todo de su mayor densidad de intersecciones. "
             "Conviene recordar que el índice es una síntesis: su lectura debe acompañarse de "
             "los indicadores individuales y del contexto de cada zona.")

    # Figura del índice, si existe
    fig = FIG_DIR / "graf_indice_caminabilidad_barrio_obrero.png"
    if fig.exists():
        doc.add_picture(str(fig), width=Inches(6.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    _parrafo(doc,
             "Nota metodológica: los umbrales de referencia se basan en literatura de "
             "caminabilidad (US EPA EnviroAtlas; IPEN / U. de Melbourne; regulación de "
             "Queensland; Stangl, 2015) y se usan como marco comparativo. El índice es "
             "ajustable y está pensado para evolucionar a medida que se incorporen nuevos "
             "indicadores o datos locales.",
             size=Pt(9.5), color=GRIS)

    doc.save(DOCX)
    print(f"Sección de índice agregada al informe: {DOCX}")


if __name__ == "__main__":
    main()
