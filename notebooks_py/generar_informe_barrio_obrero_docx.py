"""
Genera un informe en Word (.docx) de caminabilidad de Barrio Obrero.

Compila las métricas de línea base, la comparación con la Comuna 9 y las figuras
generadas por el notebook (caminabilidad_barrio_obrero_v2.ipynb) en un documento
editable, con texto justificado y una breve explicación en cada sección.

NOTA: se excluye la sección "Indicadores complementarios" (Celda 10), por no ser
relevante para el análisis de caminabilidad de este trabajo.

Uso:
    uv run notebooks_py/generar_informe_barrio_obrero_docx.py

Requiere: python-docx (uv pip install python-docx)
"""

import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

PROJECT_ROOT = Path(__file__).resolve().parent.parent
FIG_DIR = PROJECT_ROOT / "outputs" / "figures"
INFORMES_DIR = PROJECT_ROOT / "outputs" / "informes"
INFORMES_DIR.mkdir(parents=True, exist_ok=True)

FECHA = datetime.date.today().isoformat()
AZUL = RGBColor(0x1B, 0x4F, 0x8A)
VERDE = RGBColor(0x2E, 0x7D, 0x32)
GRIS = RGBColor(0x44, 0x44, 0x44)

METRICAS_BO = [
    ("Área del polígono", "37.73 ha (377,329 m²)"),
    ("Intersecciones peatonales", "233"),
    ("Longitud de red peatonal", "26.93 km"),
    ("Longitud promedio de segmento", "41.6 m"),
    ("Densidad de calle", "71.38 km/km²"),
    ("Densidad de intersecciones", "617 int/km²"),
    ("Nodos OSM", "233"),
    ("Segmentos OSM", "648"),
]

COMPARACION = [
    ("Métrica", "Barrio Obrero", "Comuna 9"),
    ("Área (ha)", "37.73", "290.38"),
    ("Intersecciones", "233", "2.255"),
    ("Longitud de red (km)", "26.93", "227.52"),
    ("Longitud prom. de segmento (m)", "41.6", "34.8"),
    ("Densidad de calle (km/km²)", "71.38", "78.35"),
    ("Densidad de intersecciones (int/km²)", "617", "777"),
]

# (título de sección, archivo de figura, texto explicativo profesional y amable)
SECCIONES_FIG = [
    ("Red peatonal — línea base",
     "graf_01_red_peatonal_barrio_obrero.png",
     "Este mapa muestra la red peatonal del barrio tal como se descarga de OpenStreetMap. "
     "Cada línea es un tramo caminable y cada punto una intersección. Nos sirve como "
     "«fotografía» inicial de cómo está conectado el barrio para quien se desplaza a pie, y "
     "es la base sobre la que se calculan todas las métricas de este informe."),
    ("Densidad de intersecciones y malla vial",
     "graf_02_heatmap_intersecciones_barrio_obrero.png",
     "Aquí presentamos un mapa de calor que resalta las zonas con mayor concentración de "
     "intersecciones. A la izquierda se ve sobre el contorno del barrio y a la derecha con "
     "los ejes viales oficiales (IDESC) de fondo. Las áreas más cálidas indican mallas más "
     "densas y, por tanto, más opciones de recorrido para el peatón."),
    ("Distribución de la longitud de segmentos",
     "graf_03_histograma_segmentos_barrio_obrero.png",
     "Este histograma resume qué tan largos son los tramos de calle del barrio. En general, "
     "los tramos cortos favorecen la caminabilidad porque ofrecen más cruces y rutas "
     "alternativas. La media y la mediana marcadas ayudan a ubicar el «tramo típico» de la zona."),
    ("Dispersión de la longitud de segmentos",
     "graf_04_boxplot_segmentos_barrio_obrero.png",
     "El diagrama de caja complementa al histograma: muestra la mediana, el rango habitual y "
     "los tramos atípicos (muy largos o muy cortos). Es útil para detectar irregularidades en "
     "la trama urbana que podrían afectar la experiencia de caminar."),
    ("Posición morfológica — Barrio Obrero vs Comuna 9",
     "graf_05_scatter_conectividad_barrio_obrero.png",
     "Este gráfico ubica al barrio según dos rasgos clave: la longitud promedio de segmento y "
     "la densidad de intersecciones, y lo compara con la Comuna 9. Las bandas de color señalan "
     "rangos de referencia de la literatura: la franja verde (400–700 int/km²) corresponde a "
     "tramas urbanas caminables y la línea roja (~100 int/km²) es el umbral asociado a "
     "beneficios de salud pública."),
    ("Tabla resumen de métricas urbanas",
     "graf_06_tabla_metricas_barrio_obrero.png",
     "Esta tabla reúne, en un solo vistazo, los indicadores más importantes de caminabilidad "
     "del barrio. Sirve como ficha técnica de la línea base y facilita el seguimiento en "
     "mediciones futuras."),
    ("Panel de métricas de caminabilidad",
     "graf_metricas_caminabilidad_barrio_obrero.png",
     "El panel presenta las cinco métricas de caminabilidad del barrio en formato de barras, "
     "pensado para una lectura rápida y para acompañar presentaciones o resúmenes ejecutivos."),
    ("Comparación Barrio Obrero vs Comuna 9",
     "graf_comparacion_barrio_obrero_vs_comuna9.png",
     "Finalmente, comparamos el barrio con la Comuna 9 que lo contiene. Conviene mirar sobre "
     "todo las densidades (km/km² e int/km²), que son comparables entre zonas de distinto "
     "tamaño; las cifras absolutas favorecen a la comuna simplemente por su mayor superficie."),
]


def _justificar(par):
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _parrafo(doc, texto, size=11, color=None, bold=False, justify=True, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if justify:
        _justificar(p)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _encabezado_seccion(doc, texto):
    h = doc.add_heading(level=1)
    run = h.add_run(texto)
    run.font.color.rgb = AZUL
    return h


def main():
    doc = Document()

    # Estilo base
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # ---- Portada ----
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Índice de Caminabilidad")
    r.bold = True; r.font.size = Pt(26); r.font.color.rgb = AZUL
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = s.add_run("Barrio Obrero — Cali")
    rs.font.size = Pt(18); rs.font.color.rgb = GRIS
    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rst = st.add_run("Informe de línea base de caminabilidad")
    rst.font.size = Pt(13); rst.bold = True; rst.font.color.rgb = VERDE
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rmeta = meta.add_run(f"ITT Cali Inteligente · Equipo de Gobierno de Datos\nFecha: {FECHA}")
    rmeta.font.size = Pt(11); rmeta.font.color.rgb = GRIS

    doc.add_paragraph()

    # ---- Presentación ----
    _encabezado_seccion(doc, "Presentación")
    _parrafo(doc,
             "Este informe resume el estado de la caminabilidad del Barrio Obrero (Cali) a "
             "partir de la red peatonal de OpenStreetMap. Nuestro objetivo es ofrecer una "
             "línea base clara y fácil de leer, que sirva como punto de partida para evaluar "
             "el impacto de futuras intervenciones y para dialogar con distintos equipos del "
             "programa. Hemos procurado un lenguaje sencillo, acompañando cada sección con una "
             "breve explicación de lo que muestra y de por qué es relevante.")
    _parrafo(doc,
             "La unidad de análisis es el polígono único del barrio (37.73 ha). Todo el trabajo "
             "geoespacial se realiza en WGS84 (EPSG:4326) y se emplea EPSG:3116 únicamente para "
             "los cálculos de área y distancia. Con el fin de mantener el foco en la "
             "caminabilidad, se excluye deliberadamente la sección de indicadores "
             "complementarios (eventos), que no resulta relevante para este análisis.")

    # ---- Metodología ----
    _encabezado_seccion(doc, "Metodología en breve")
    _parrafo(doc,
             "La red peatonal se descarga con la librería OSMnx (tipo de red «walk») "
             "conservando todos los vértices de cada calle. Sobre ese grafo se calculan las "
             "métricas de conectividad: número de intersecciones, longitud total de la red, "
             "longitud promedio de los tramos y densidades por unidad de área. Las densidades "
             "son especialmente útiles porque permiten comparar zonas de distinto tamaño en "
             "igualdad de condiciones.")

    # ---- Métricas de línea base ----
    _encabezado_seccion(doc, "Métricas de línea base")
    _parrafo(doc,
             "La siguiente tabla reúne los valores de referencia del barrio. Son las cifras que "
             "usaremos como punto de comparación en mediciones posteriores.")
    tabla = doc.add_table(rows=1, cols=2)
    tabla.style = "Light Grid Accent 1"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tabla.rows[0].cells
    hdr[0].paragraphs[0].add_run("Indicador").bold = True
    hdr[1].paragraphs[0].add_run("Valor").bold = True
    for k, v in METRICAS_BO:
        row = tabla.add_row().cells
        row[0].text = k
        row[1].text = v

    # ---- Comparación con Comuna 9 ----
    _encabezado_seccion(doc, "Contexto territorial: comparación con la Comuna 9")
    _parrafo(doc,
             "Para entender mejor el barrio, lo comparamos con la Comuna 9, que lo contiene. "
             "Recomendamos leer sobre todo las densidades, ya que las cifras absolutas "
             "(intersecciones y longitud total) favorecen a la comuna simplemente por su mayor "
             "extensión.")
    tabla2 = doc.add_table(rows=1, cols=3)
    tabla2.style = "Light Grid Accent 1"
    tabla2.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr2 = tabla2.rows[0].cells
    for j, titulo in enumerate(COMPARACION[0]):
        hdr2[j].paragraphs[0].add_run(titulo).bold = True
    for fila in COMPARACION[1:]:
        row = tabla2.add_row().cells
        for j, val in enumerate(fila):
            row[j].text = val
    _parrafo(doc,
             "En síntesis, el Barrio Obrero presenta una densidad de intersecciones de "
             "617 int/km² y una densidad de calle de 71.38 km/km², valores propios de una trama "
             "urbana consolidada y caminable, muy por encima del umbral de salud pública "
             "(~100 int/km²). La Comuna 9 es algo más densa (777 int/km²) y con tramos más "
             "cortos, lo que ubica al barrio como una zona bien conectada dentro de un contexto "
             "comunal favorable para caminar.")

    # ---- Resultados gráficos ----
    _encabezado_seccion(doc, "Resultados gráficos")
    _parrafo(doc,
             "A continuación se presentan los gráficos de la línea base. Cada uno incluye una "
             "breve nota que explica qué observar y por qué es importante.")
    for titulo, archivo, explicacion in SECCIONES_FIG:
        h = doc.add_heading(level=2)
        rh = h.add_run(titulo); rh.font.color.rgb = VERDE
        _parrafo(doc, explicacion)
        ruta = FIG_DIR / archivo
        if ruta.exists():
            doc.add_picture(str(ruta), width=Inches(6.2))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            _parrafo(doc, f"[Figura no encontrada: {archivo}]", color=GRIS, justify=False)

    # ---- Referencias ----
    _encabezado_seccion(doc, "Referencias de contexto")
    for ref in [
        "US EPA — Estimated Intersection Density of Walkable Roads (EnviroAtlas).",
        "IPEN Adult study / Universidad de Melbourne — umbral de ~100 intersecciones/km² "
        "asociado a mayor actividad física.",
        "Gobierno de Queensland — regulación de barrios caminables (cuadras ≤ 250 m).",
        "Stangl, P. (2015) — Block size-based measures of street connectivity, Urban Design "
        "International.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(ref)
        _justificar(p)
    _parrafo(doc,
             "Nota: los umbrales provienen mayormente de literatura anglosajona (EE. UU. y "
             "Australia). No existe un estándar nacional colombiano de densidad de "
             "intersecciones ampliamente adoptado, por lo que aquí se usan como referencia "
             "comparativa, con la Comuna 9 aportando el contraste local. Contenido reformulado "
             "para cumplir con las restricciones de licencia de las fuentes.",
             size=9.5, color=GRIS)

    salida = INFORMES_DIR / "informe_caminabilidad_barrio_obrero.docx"
    doc.save(salida)
    print(f"Informe Word generado: {salida}")


if __name__ == "__main__":
    main()
